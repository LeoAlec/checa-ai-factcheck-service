"""Reproducible, leakage-aware import of the local FakeRecogna workbook."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook

from .contracts import DatasetSplit, TruthLabel
from .fakewhatsapp import _normalize_text, _stratified_sample

_DATE = re.compile(r"(\d{1,2}/\d{1,2}/\d{4})")


@dataclass(frozen=True)
class FakeRecognaImportSummary:
    raw_rows: int
    labeled_rows: int
    excluded_missing_fields: int
    excluded_non_target_domain: int
    excluded_outside_time_window: int
    excluded_long_rows: int
    duplicate_rows: int
    conflicting_duplicate_rows: int
    calibration_candidates: int
    test_candidates: int
    calibration_exported: int
    test_exported: int
    seed: int

    def as_dict(self) -> dict[str, object]:
        return asdict(self)


def _parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    match = _DATE.search(str(value or ""))
    if not match:
        return None
    try:
        return datetime.strptime(match.group(1), "%d/%m/%Y").date()
    except ValueError:
        return None


def _domain(value: str) -> str:
    return urlparse(value).netloc.casefold().removeprefix("www.")


def build_fakerecogna_dataset(
    input_path: str | Path,
    output_path: str | Path,
    *,
    manifest_path: str | Path | None = None,
    domains: tuple[str, ...] = ("g1.globo.com",),
    calibration_until: date = date(2020, 12, 31),
    test_from: date = date(2021, 1, 1),
    max_calibration: int = 750,
    max_test: int = 400,
    max_text_chars: int = 5_000,
    seed: int = 42,
) -> FakeRecognaImportSummary:
    """Export title-only, date-separated records from existing workbook labels.

    ``Classe=0`` is the false class and ``Classe=1`` the true class.  The default
    limits selection to G1, a source containing both labels, to avoid an evaluation
    dominated by source identity (many other workbook domains occur in one class).
    The body column is deliberately never used because it contains processed news
    text and can leak a fact-check explanation into the classifier input.
    """
    if calibration_until >= test_from:
        raise ValueError("calibration_until deve ser anterior a test_from")
    if max_calibration < 2 or max_test < 2:
        raise ValueError("cada split deve permitir ao menos duas mensagens")
    accepted_domains = {item.casefold().removeprefix("www.") for item in domains}
    if not accepted_domains:
        raise ValueError("informe ao menos um domínio")

    sheet = load_workbook(Path(input_path), read_only=True, data_only=True).active
    header = next(sheet.iter_rows(values_only=True))
    positions = {str(name): index for index, name in enumerate(header)}
    required = {"Titulo", "Categoria", "Data", "URL", "Classe"}
    missing = required.difference(positions)
    if missing:
        raise ValueError(f"colunas ausentes no FakeRecogna: {', '.join(sorted(missing))}")

    candidates: list[dict[str, str]] = []
    raw_rows = labeled_rows = missing_fields = non_target = outside = long_rows = 0
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 2):
        raw_rows += 1
        label_value = row[positions["Classe"]]
        if label_value not in (0, 1):
            continue
        labeled_rows += 1
        title = str(row[positions["Titulo"]] or "").strip()
        source = str(row[positions["URL"]] or "").strip()
        topic = str(row[positions["Categoria"]] or "").strip()
        published_at = _parse_date(row[positions["Data"]])
        if not title or not source or not topic or published_at is None:
            missing_fields += 1
            continue
        if len(title) > max_text_chars:
            long_rows += 1
            continue
        domain = _domain(source)
        if domain not in accepted_domains:
            non_target += 1
            continue
        if published_at <= calibration_until:
            split = DatasetSplit.CALIBRATION
        elif published_at >= test_from:
            split = DatasetSplit.TEST
        else:
            outside += 1
            continue
        candidates.append({
            "id": f"fakerecogna-{row_number}", "text": title,
            "label": TruthLabel.FALSE.value if label_value == 0 else TruthLabel.TRUE.value,
            "topic": topic, "source": source, "published_at": published_at.isoformat(),
            "split": split.value, "_normalized_text": _normalize_text(title),
        })

    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in candidates:
        grouped[row["_normalized_text"]].append(row)
    deduplicated: list[dict[str, str]] = []
    duplicate_rows = conflicting_rows = 0
    for duplicates in grouped.values():
        labels = {row["label"] for row in duplicates}
        if len(duplicates) == 1:
            deduplicated.extend(duplicates)
        elif len(labels) == 1:
            deduplicated.append(sorted(duplicates, key=lambda row: row["id"])[0])
            duplicate_rows += len(duplicates) - 1
        else:
            conflicting_rows += len(duplicates)
    calibration = [row for row in deduplicated if row["split"] == "calibration"]
    test = [row for row in deduplicated if row["split"] == "test"]
    selected = [*_stratified_sample(calibration, max_calibration, seed), *_stratified_sample(test, max_test, seed + 1)]
    for row in selected:
        row.pop("_normalized_text", None)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=("id", "text", "label", "topic", "source", "published_at", "split"))
        writer.writeheader()
        writer.writerows(selected)
    summary = FakeRecognaImportSummary(raw_rows, labeled_rows, missing_fields, non_target, outside, long_rows,
                                       duplicate_rows, conflicting_rows, len(calibration), len(test),
                                       sum(row["split"] == "calibration" for row in selected),
                                       sum(row["split"] == "test" for row in selected), seed)
    if manifest_path:
        manifest = {
            "dataset": "FakeRecogna", "input": str(input_path), "output": str(output),
            "input_field": "Titulo (somente; Noticia é excluída para evitar vazamento)",
            "label_mapping": {"0": "falso", "1": "verdadeiro"}, "accepted_domains": sorted(accepted_domains),
            "split": {"calibration_until": calibration_until.isoformat(), "test_from": test_from.isoformat()},
            "limitations": [
                "a origem do workbook e possível sobreposição com dados de treino precisam de auditoria independente",
                "a seleção G1 reduz, mas não elimina, vieses por fonte e por estilo editorial",
                "os rótulos são reaproveitados do workbook; nenhuma alegação foi re-rotulada por LLM",
            ], "summary": summary.as_dict(),
            "input_sha256": hashlib.sha256(Path(input_path).read_bytes()).hexdigest(),
        }
        Path(manifest_path).write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary
